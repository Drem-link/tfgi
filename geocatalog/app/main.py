import csv
import hashlib
import io
import json
import math
import os
import re
import secrets
import time
import zipfile
from contextlib import asynccontextmanager, closing
from datetime import date, datetime
from pathlib import Path
from typing import Annotated
from urllib.parse import parse_qs, urlsplit
from xml.etree.ElementTree import ParseError

import psycopg
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from psycopg.errors import UniqueViolation
from psycopg.rows import dict_row
from argon2 import PasswordHasher
from argon2.exceptions import InvalidHashError, VerificationError, VerifyMismatchError
from fastapi import FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, model_validator
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer

password_hasher = PasswordHasher()
login_failures: dict[str, list[float]] = {}
session_max_age = 8 * 60 * 60
login_window = 15 * 60
login_failure_limit = 8
max_file_size = 50 * 1024 * 1024
max_upload_files = 10
max_request_size = 100 * 1024 * 1024
max_csv_import_size = 5 * 1024 * 1024
max_csv_import_rows = 1000
upload_root = Path(os.environ.get("UPLOAD_DIR", "/uploads"))
allowed_file_types = {
    ".pdf": "application/pdf",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".txt": "text/plain",
    ".csv": "text/csv",
}


def auth_config() -> tuple[str, str, URLSafeTimedSerializer]:
    username = os.environ.get("AUTH_USERNAME", "").strip()
    password_hash = os.environ.get("AUTH_PASSWORD_HASH", "").strip()
    session_secret = os.environ.get("SESSION_SECRET", "").strip()
    if not username or not password_hash or len(session_secret) < 32:
        raise RuntimeError("AUTH_USERNAME, AUTH_PASSWORD_HASH and a 32-character SESSION_SECRET are required")
    return username, password_hash, URLSafeTimedSerializer(session_secret, salt="geocatalog-session-v1")


def ensure_users_table() -> None:
    username, password_hash, _ = auth_config()
    with get_connection() as connection:
        connection.execute("SELECT pg_advisory_xact_lock(73412906)")
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS geocatalog_users (
                id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
                username TEXT NOT NULL UNIQUE
                    CHECK (username ~ '^[A-Za-z0-9._@-]{1,128}$'),
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL CHECK (role IN ('admin', 'viewer')),
                is_active BOOLEAN NOT NULL DEFAULT TRUE,
                token_version INTEGER NOT NULL DEFAULT 0,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """
        )
        connection.execute(
            """
            INSERT INTO geocatalog_users (username, password_hash, role)
            VALUES (%s, %s, 'admin')
            ON CONFLICT (username) DO NOTHING
            """,
            (username, password_hash),
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS document_files (
                id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
                document_id BIGINT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                storage_key TEXT NOT NULL UNIQUE CHECK (storage_key ~ '^[a-f0-9]{48}\\.[a-z0-9]+$'),
                original_filename TEXT NOT NULL,
                media_type TEXT NOT NULL,
                size_bytes BIGINT NOT NULL CHECK (size_bytes BETWEEN 1 AND 52428800),
                created_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """
        )
        connection.execute("CREATE INDEX IF NOT EXISTS document_files_document_idx ON document_files (document_id)")
        connection.execute("ALTER TABLE features ALTER COLUMN geom DROP NOT NULL")
        connection.execute("ALTER TABLE features ADD COLUMN IF NOT EXISTS source_geom geometry(Geometry)")
        connection.execute("ALTER TABLE features ADD COLUMN IF NOT EXISTS source_crs TEXT NOT NULL DEFAULT 'EPSG:4326'")
        connection.execute(
            "UPDATE features SET source_geom = geom WHERE source_geom IS NULL AND geom IS NOT NULL"
        )
        for column, definition in (
            ("tgf_number", "TEXT"),
            ("document_type", "TEXT"),
            ("authors", "TEXT"),
            ("coauthors", "TEXT"),
            ("executor_org", "TEXT"),
            ("work_year_start", "INTEGER"),
            ("work_year_end", "INTEGER"),
            ("created_place", "TEXT"),
            ("minerals", "TEXT"),
            ("archive_disk_number", "TEXT"),
            ("material_composition", "TEXT"),
            ("electronic_copy_status", "TEXT"),
            ("efgi_id", "TEXT"),
            ("efgi_url", "TEXT"),
        ):
            connection.execute(f"ALTER TABLE documents ADD COLUMN IF NOT EXISTS {column} {definition}")
        connection.execute("ALTER TABLE documents ADD COLUMN IF NOT EXISTS import_fingerprint CHAR(64)")
        connection.execute(
            "CREATE UNIQUE INDEX IF NOT EXISTS documents_import_fingerprint_idx "
            "ON documents (import_fingerprint) WHERE import_fingerprint IS NOT NULL"
        )
        connection.execute(
            """
            DO $$ BEGIN
                IF NOT EXISTS (
                    SELECT 1 FROM pg_constraint
                    WHERE conname = 'features_source_crs_check'
                      AND conrelid = 'features'::regclass
                ) THEN
                    ALTER TABLE features ADD CONSTRAINT features_source_crs_check CHECK (
                        source_crs IN ('EPSG:4326', 'EPSG:7683')
                        AND (source_geom IS NULL OR ST_SRID(source_geom) =
                            CASE source_crs WHEN 'EPSG:7683' THEN 7683 ELSE 4326 END)
                        AND (source_crs <> 'EPSG:7683' OR geom IS NULL)
                    );
                END IF;
                IF NOT EXISTS (
                    SELECT 1 FROM pg_constraint
                    WHERE conname = 'documents_work_years_check'
                      AND conrelid = 'documents'::regclass
                ) THEN
                    ALTER TABLE documents ADD CONSTRAINT documents_work_years_check CHECK (
                        work_year_start IS NULL OR work_year_end IS NULL OR work_year_start <= work_year_end
                    );
                END IF;
                IF NOT EXISTS (
                    SELECT 1 FROM pg_constraint
                    WHERE conname = 'documents_efgi_url_check'
                      AND conrelid = 'documents'::regclass
                ) THEN
                    ALTER TABLE documents ADD CONSTRAINT documents_efgi_url_check CHECK (
                        efgi_url IS NULL OR efgi_url ~ '^https?://[^[:space:]]+$'
                    );
                END IF;
            END $$;
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS document_relations (
                id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
                source_document_id BIGINT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                target_document_id BIGINT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                relation_type TEXT NOT NULL CHECK (
                    relation_type IN ('related', 'appendix', 'protocol', 'map', 'copy_of')
                ),
                created_by BIGINT REFERENCES geocatalog_users(id) ON DELETE SET NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                CHECK (source_document_id <> target_document_id),
                UNIQUE (source_document_id, target_document_id, relation_type)
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS saved_searches (
                id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
                user_id BIGINT NOT NULL REFERENCES geocatalog_users(id) ON DELETE CASCADE,
                name TEXT NOT NULL CHECK (length(name) BETWEEN 1 AND 120),
                filters JSONB NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                UNIQUE (user_id, name)
            )
            """
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS document_relations_target_idx ON document_relations (target_document_id)"
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS catalog_audit (
                id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
                actor_id BIGINT REFERENCES geocatalog_users(id) ON DELETE SET NULL,
                action TEXT NOT NULL,
                entity_type TEXT NOT NULL,
                entity_id BIGINT,
                details JSONB NOT NULL DEFAULT '{}'::jsonb,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS catalog_audit_created_idx ON catalog_audit (created_at DESC)"
        )


@asynccontextmanager
async def app_lifespan(_: FastAPI):
    ensure_users_table()
    yield


app = FastAPI(title="Geological Library Map", version="0.1.0", lifespan=app_lifespan)
app.mount("/static", StaticFiles(directory="app/static"), name="static")


def authenticated_user(request: Request) -> dict | None:
    try:
        _, _, serializer = auth_config()
        session = serializer.loads(
            request.cookies.get("geocatalog_session", ""),
            max_age=session_max_age,
        )
    except (RuntimeError, BadSignature, SignatureExpired, TypeError):
        return None
    if not isinstance(session, dict) or not isinstance(session.get("id"), int):
        return None
    with closing(get_connection()) as connection:
        user = connection.execute(
            """
            SELECT id, username, role
            FROM geocatalog_users
            WHERE id = %s AND token_version = %s AND is_active
            """,
            (session["id"], session.get("version")),
        ).fetchone()
    if user is None:
        return None
    return user


def is_public_path(path: str) -> bool:
    return path in {
        "/login",
        "/auth/login",
        "/healthz",
        "/livez",
        "/static/login.css",
        "/static/login.js",
        "/static/ne_110m_admin_0_countries.geojson",
    }


@app.middleware("http")
async def require_login(request: Request, call_next):
    is_upload = request.url.path == "/api/admin/documents" and request.method == "POST"
    is_bulk_import = request.url.path.startswith("/api/admin/import/") and request.method == "POST"
    if is_public_path(request.url.path):
        response = await call_next(request)
    else:
        user = authenticated_user(request)
        if user is None:
            if request.url.path.startswith("/api/"):
                response = JSONResponse({"detail": "authentication required"}, status_code=401)
            else:
                response = RedirectResponse("/login", status_code=303)
        else:
            request.state.user = user
            if is_upload and user["role"] != "admin":
                response = JSONResponse({"detail": "administrator permission required"}, status_code=403)
            elif is_upload and request.headers.get("content-length") is None:
                response = JSONResponse({"detail": "Content-Length is required for file uploads"}, status_code=411)
            elif is_upload and not request.headers.get("content-length", "").isdigit():
                response = JSONResponse({"detail": "invalid Content-Length"}, status_code=400)
            elif is_upload and int(request.headers["content-length"]) > max_request_size + 5 * 1024 * 1024:
                response = JSONResponse({"detail": "upload request cannot exceed 105 MiB"}, status_code=413)
            elif is_bulk_import and request.headers.get("content-length") is None:
                response = JSONResponse({"detail": "Content-Length is required for import"}, status_code=411)
            elif is_bulk_import and not request.headers.get("content-length", "").isdigit():
                response = JSONResponse({"detail": "invalid Content-Length"}, status_code=400)
            elif is_bulk_import and int(request.headers["content-length"]) > max_csv_import_size:
                response = JSONResponse({"detail": "import file cannot exceed 5 MiB"}, status_code=413)
            elif request.url.path.startswith("/api/admin/") and user["role"] != "admin":
                response = JSONResponse({"detail": "administrator permission required"}, status_code=403)
            elif request.url.path.startswith("/admin") and user["role"] != "admin":
                response = RedirectResponse("/", status_code=303)
            else:
                response = await call_next(request)
    if request.url.scheme == "https":
        response.headers["Strict-Transport-Security"] = "max-age=31536000"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["Cache-Control"] = "no-store"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; script-src 'self'; style-src 'self'; "
        "img-src 'self' data:; connect-src 'self'; font-src 'self' data:; "
        "object-src 'none'; base-uri 'self'; frame-ancestors 'none'; form-action 'self'"
    )
    return response


def get_connection():
    database_url = os.environ.get("DATABASE_URL")
    if database_url:
        return psycopg.connect(database_url, row_factory=dict_row)
    required = ("PGHOST", "PGDATABASE", "PGUSER", "PGPASSWORD")
    missing = [key for key in required if not os.environ.get(key)]
    if missing:
        raise RuntimeError(f"Database connection settings are missing: {', '.join(missing)}")
    return psycopg.connect(
        host=os.environ["PGHOST"],
        port=os.environ.get("PGPORT", "5432"),
        dbname=os.environ["PGDATABASE"],
        user=os.environ["PGUSER"],
        password=os.environ["PGPASSWORD"],
        connect_timeout=5,
        row_factory=dict_row,
    )


class UserCreate(BaseModel):
    username: str = Field(min_length=1, max_length=128, pattern=r"^[A-Za-z0-9._@-]+$")
    password: str = Field(min_length=8, max_length=1024)
    role: str = Field(pattern=r"^(admin|viewer)$")


class UserUpdate(BaseModel):
    role: str | None = Field(default=None, pattern=r"^(admin|viewer)$")
    active: bool | None = None
    password: str | None = Field(default=None, min_length=8, max_length=1024)

    @model_validator(mode="after")
    def require_update(self):
        if self.role is None and self.active is None and self.password is None:
            raise ValueError("at least one of role, active, or password is required")
        return self


def require_same_origin(request: Request) -> None:
    origin = request.headers.get("origin")
    if origin and urlsplit(origin).netloc != request.headers.get("host", ""):
        raise HTTPException(status_code=403, detail="invalid request origin")
    if not origin and request.headers.get("sec-fetch-site") == "cross-site":
        raise HTTPException(status_code=403, detail="cross-site request denied")


def parse_coordinates(value: str) -> tuple[float, float]:
    try:
        longitude, latitude = map(float, value.split(","))
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="each vertex must be longitude,latitude") from exc
    if not math.isfinite(longitude) or not math.isfinite(latitude):
        raise HTTPException(status_code=422, detail="coordinates must be finite numbers")
    if not -180 <= longitude <= 180 or not -90 <= latitude <= 90:
        raise HTTPException(status_code=422, detail="coordinates are outside longitude/latitude ranges")
    return longitude, latitude


def is_http_url(value: str) -> bool:
    try:
        parsed = urlsplit(value)
        return (
            not any(character.isspace() for character in value)
            and parsed.scheme in {"http", "https"}
            and parsed.hostname is not None
            and parsed.username is None
            and parsed.password is None
        )
    except ValueError:
        return False


def csv_export_value(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "\t" + value
    return value


def import_fingerprint(row: dict) -> str:
    canonical = {key: value for key, value in row.items() if key not in {"row", "errors"}}
    serialized = json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def parse_polygon_vertices(value: str) -> list[tuple[float, float]]:
    if len(value) > 20000:
        raise HTTPException(status_code=422, detail="an area contour cannot exceed 20,000 characters")
    vertices = [parse_coordinates(pair.strip()) for pair in value.split(";") if pair.strip()]
    if len(vertices) < 3:
        raise HTTPException(status_code=422, detail="an area needs at least three vertices")
    if len(vertices) > 2000:
        raise HTTPException(status_code=422, detail="an area cannot exceed 2,000 vertices")
    if vertices[-1] != vertices[0]:
        vertices.append(vertices[0])
    if len(set(vertices[:-1])) < 3:
        raise HTTPException(status_code=422, detail="an area needs at least three distinct vertices")
    return vertices


def validate_uploaded_file(filename: str | None, data: bytes) -> tuple[str, str, str]:
    if not filename:
        raise HTTPException(status_code=422, detail="every uploaded file must have a filename")
    normalized_name = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not normalized_name or len(normalized_name) > 255 or any(ord(char) < 32 for char in normalized_name):
        raise HTTPException(status_code=422, detail="invalid filename")
    extension = Path(normalized_name).suffix.lower()
    media_type = allowed_file_types.get(extension)
    if media_type is None:
        raise HTTPException(status_code=415, detail=f"file type {extension or '(no extension)'} is not allowed")
    if not data or len(data) > max_file_size:
        raise HTTPException(status_code=413, detail="each file must be between 1 byte and 50 MiB")

    valid = False
    if extension == ".pdf":
        valid = data.startswith(b"%PDF-")
    elif extension in {".jpg", ".jpeg"}:
        valid = data.startswith(b"\xff\xd8\xff")
    elif extension == ".png":
        valid = data.startswith(b"\x89PNG\r\n\x1a\n")
    elif extension in {".tif", ".tiff"}:
        valid = data.startswith((b"II*\x00", b"MM\x00*"))
    elif extension in {".docx", ".xlsx", ".pptx"}:
        expected_part = {
            ".docx": "word/",
            ".xlsx": "xl/",
            ".pptx": "ppt/",
        }[extension]
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
                valid = "[Content_Types].xml" in names and any(name.startswith(expected_part) for name in names)
        except (zipfile.BadZipFile, OSError):
            valid = False
    elif extension in {".txt", ".csv"}:
        try:
            text_data = data.decode("utf-8-sig")
            valid = "\x00" not in text_data
            if valid and extension == ".csv":
                next(csv.reader(io.StringIO(text_data)), None)
        except (UnicodeDecodeError, csv.Error):
            valid = False

    if not valid:
        raise HTTPException(status_code=415, detail=f"file contents do not match {extension}")
    return normalized_name, media_type, extension[1:]


def require_admin(request: Request) -> dict:
    user = getattr(request.state, "user", None)
    if user is None or user["role"] != "admin":
        raise HTTPException(status_code=403, detail="administrator permission required")
    return user


def parse_bbox(bbox: str | None) -> tuple[float, float, float, float] | None:
    if bbox is None:
        return None
    try:
        west, south, east, north = map(float, bbox.split(","))
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="bbox must be west,south,east,north") from exc
    if (
        not all(map(math.isfinite, (west, south, east, north)))
        or not all(-180 <= value <= 180 for value in (west, east))
        or not all(-90 <= value <= 90 for value in (south, north))
        or west >= east
        or south >= north
    ):
        raise HTTPException(status_code=422, detail="bbox coordinates are out of range or reversed")
    return west, south, east, north


def decode_geojson(value):
    return json.loads(value) if isinstance(value, str) else value


def attach_document_files(connection, feature_rows) -> None:
    documents = [document for row in feature_rows for document in row["documents"]]
    document_ids = [document["id"] for document in documents]
    if not document_ids:
        return
    files = connection.execute(
        """
        SELECT id, document_id, original_filename, size_bytes
        FROM document_files
        WHERE document_id = ANY(%s)
        ORDER BY original_filename
        """,
        (document_ids,),
    ).fetchall()
    files_by_document: dict[int, list[dict]] = {}
    for file in files:
        files_by_document.setdefault(file["document_id"], []).append({
            "id": file["id"],
            "filename": file["original_filename"],
            "size_bytes": file["size_bytes"],
        })
    for document in documents:
        document["files"] = files_by_document.get(document["id"], [])
    relations = connection.execute(
        """
        SELECT doc_relation.id, doc_relation.source_document_id, doc_relation.target_document_id,
               doc_relation.relation_type, source_doc.title AS source_title, target_doc.title AS target_title
        FROM document_relations doc_relation
        JOIN documents source_doc ON source_doc.id = doc_relation.source_document_id
        JOIN documents target_doc ON target_doc.id = doc_relation.target_document_id
        WHERE doc_relation.source_document_id = ANY(%s) OR doc_relation.target_document_id = ANY(%s)
        ORDER BY doc_relation.id
        """,
        (document_ids, document_ids),
    ).fetchall()
    related_by_document: dict[int, list[dict]] = {}
    for relation in relations:
        source_id = relation["source_document_id"]
        target_id = relation["target_document_id"]
        related_by_document.setdefault(source_id, []).append({
            "id": relation["id"], "document_id": target_id,
            "title": relation["target_title"], "relation_type": relation["relation_type"],
        })
        related_by_document.setdefault(target_id, []).append({
            "id": relation["id"], "document_id": source_id,
            "title": relation["source_title"], "relation_type": relation["relation_type"],
        })
    for document in documents:
        document["relations"] = related_by_document.get(document["id"], [])


def write_audit(connection, user_id: int, action: str, entity_type: str, entity_id: int | None, details: dict) -> None:
    connection.execute(
        """
        INSERT INTO catalog_audit (actor_id, action, entity_type, entity_id, details)
        VALUES (%s, %s, %s, %s, %s)
        """,
        (user_id, action, entity_type, entity_id, json.dumps(details, ensure_ascii=False)),
    )


def document_json_sql(alias: str = "d") -> str:
    return f"""json_build_object(
        'id', {alias}.id,
        'title', {alias}.title,
        'inventory_number', {alias}.inventory_number,
        'region', {alias}.region,
        'year', {alias}.year,
        'topic', {alias}.topic,
        'description', {alias}.description,
        'archive_reference', {alias}.archive_reference,
        'tgf_number', {alias}.tgf_number,
        'document_type', {alias}.document_type,
        'authors', {alias}.authors,
        'coauthors', {alias}.coauthors,
        'executor_org', {alias}.executor_org,
        'work_year_start', {alias}.work_year_start,
        'work_year_end', {alias}.work_year_end,
        'created_place', {alias}.created_place,
        'minerals', {alias}.minerals,
        'archive_disk_number', {alias}.archive_disk_number,
        'material_composition', {alias}.material_composition,
        'electronic_copy_status', {alias}.electronic_copy_status,
        'efgi_id', {alias}.efgi_id,
        'efgi_url', {alias}.efgi_url
    )"""


class SavedSearchCreate(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    filters: dict[str, str | int] = Field(max_length=20)

    @model_validator(mode="after")
    def validate_filters(self):
        self.name = self.name.strip()
        if not self.name:
            raise ValueError("name cannot be blank")
        allowed = {"q", "region", "year_from", "year_to"}
        unknown = set(self.filters) - allowed
        if unknown:
            raise ValueError(f"unsupported filters: {', '.join(sorted(unknown))}")
        for key in ("q", "region"):
            value = self.filters.get(key)
            if value is not None and (not isinstance(value, str) or len(value) > 200):
                raise ValueError(f"{key} must be a string of at most 200 characters")
        for key in ("year_from", "year_to"):
            value = self.filters.get(key)
            if value is not None:
                try:
                    year = int(value)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"{key} must be a year from 1500 to 2200") from exc
                if not 1500 <= year <= 2200:
                    raise ValueError(f"{key} must be a year from 1500 to 2200")
                self.filters[key] = year
        if (
            self.filters.get("year_from") is not None
            and self.filters.get("year_to") is not None
            and self.filters["year_from"] > self.filters["year_to"]
        ):
            raise ValueError("year_from must not exceed year_to")
        return self


class DocumentUpdate(BaseModel):
    title: str | None = Field(default=None, min_length=1, max_length=500)
    inventory_number: str | None = Field(default=None, max_length=200)
    region: str | None = Field(default=None, max_length=200)
    year: int | None = Field(default=None, ge=1500, le=2200)
    topic: str | None = Field(default=None, max_length=1000)
    description: str | None = Field(default=None, max_length=10000)
    archive_reference: str | None = Field(default=None, max_length=1000)
    tgf_number: str | None = Field(default=None, max_length=200)
    document_type: str | None = Field(default=None, max_length=200)
    authors: str | None = Field(default=None, max_length=2000)
    coauthors: str | None = Field(default=None, max_length=2000)
    executor_org: str | None = Field(default=None, max_length=500)
    work_year_start: int | None = Field(default=None, ge=1500, le=2200)
    work_year_end: int | None = Field(default=None, ge=1500, le=2200)
    created_place: str | None = Field(default=None, max_length=500)
    minerals: str | None = Field(default=None, max_length=2000)
    archive_disk_number: str | None = Field(default=None, max_length=200)
    material_composition: str | None = Field(default=None, max_length=4000)
    electronic_copy_status: str | None = Field(default=None, max_length=200)
    efgi_id: str | None = Field(default=None, max_length=200)
    efgi_url: str | None = Field(default=None, max_length=2000)

    @model_validator(mode="after")
    def require_update(self):
        if not self.model_fields_set:
            raise ValueError("at least one field is required")
        if self.work_year_start and self.work_year_end and self.work_year_start > self.work_year_end:
            raise ValueError("work_year_start must not exceed work_year_end")
        if self.efgi_url and not is_http_url(self.efgi_url):
            raise ValueError("efgi_url must use HTTP or HTTPS")
        return self


class DocumentRelationCreate(BaseModel):
    target_document_id: int = Field(gt=0)
    relation_type: str = Field(pattern=r"^(related|appendix|protocol|map|copy_of)$")


def parse_catalog_csv(raw: bytes) -> list[dict]:
    if len(raw) > max_csv_import_size:
        raise HTTPException(status_code=413, detail="CSV import cannot exceed 5 MiB")
    try:
        text = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text, newline=""))
        if not reader.fieldnames:
            raise ValueError("CSV header is missing")
        reader.fieldnames = [name.strip() for name in reader.fieldnames]
        if len(reader.fieldnames) > 40 or len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError("CSV headers must be unique and cannot exceed 40 columns")
        if any(not name for name in reader.fieldnames):
            raise ValueError("CSV headers cannot be blank")
        supported = {
            "feature_name", "geometry_kind", "coordinate_crs", "document_title", "longitude", "latitude",
            "polygon_vertices", "source_wkt", "inventory_number", "tgf_number", "region", "year", "topic",
            "description", "archive_reference", "document_type", "authors", "coauthors", "executor_org",
            "work_year_start", "work_year_end", "created_place", "minerals", "archive_disk_number",
            "material_composition", "electronic_copy_status", "efgi_id", "efgi_url",
        }
        unknown = set(reader.fieldnames) - supported
        if unknown:
            raise ValueError(f"unsupported CSV columns: {', '.join(sorted(unknown))}")
        required = {"feature_name", "geometry_kind", "coordinate_crs", "document_title"}
        if not required.issubset(reader.fieldnames):
            raise ValueError(f"CSV must include columns: {', '.join(sorted(required))}")
        rows = list(reader)
    except (UnicodeDecodeError, csv.Error, ValueError) as exc:
        raise HTTPException(status_code=422, detail=f"invalid CSV: {exc}") from exc
    if not rows or len(rows) > max_csv_import_rows:
        raise HTTPException(status_code=422, detail=f"CSV must contain 1 to {max_csv_import_rows} data rows")
    validated = []
    allowed_crs = {"EPSG:7683", "EPSG:4326"}
    for index, row in enumerate(rows, start=2):
        errors = []
        if None in row:
            errors.append("row has more values than the CSV header")
        for key, value in list(row.items()):
            if isinstance(value, str) and value.startswith("\t") and value[1:].lstrip().startswith(("=", "+", "-", "@")):
                row[key] = value[1:]
        feature_name = (row.get("feature_name") or "").strip()
        title = (row.get("document_title") or "").strip()
        kind = (row.get("geometry_kind") or "").strip()
        crs = (row.get("coordinate_crs") or "").strip().upper()
        if not feature_name or len(feature_name) > 500:
            errors.append("feature_name is required (max 500)")
        if not title or len(title) > 500:
            errors.append("document_title is required (max 500)")
        if kind not in {"well", "area", "site", "other"}:
            errors.append("geometry_kind must be well, area, site or other")
        if crs not in allowed_crs:
            errors.append("coordinate_crs must be EPSG:7683 or EPSG:4326")
        wkt = None
        try:
            source_wkt = (row.get("source_wkt") or "").strip()
            if source_wkt:
                point_match = re.fullmatch(
                    r"POINT\s*\(\s*([-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][-+]?\d+)?)\s+"
                    r"([-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][-+]?\d+)?)\s*\)",
                    source_wkt,
                    re.IGNORECASE,
                )
                polygon_match = re.fullmatch(r"POLYGON\s*\(\s*\(([^()]*)\)\s*\)", source_wkt, re.IGNORECASE)
                if kind != "area" and point_match:
                    longitude, latitude = parse_coordinates(f"{point_match.group(1)},{point_match.group(2)}")
                    wkt = f"POINT({longitude} {latitude})"
                elif kind in {"area", "site", "other"} and polygon_match:
                    polygon_coords = [
                        ",".join(pair.strip().split())
                        for pair in polygon_match.group(1).split(",")
                    ]
                    vertices = parse_polygon_vertices(";".join(polygon_coords))
                    wkt = "POLYGON((" + ", ".join(f"{lon} {lat}" for lon, lat in vertices) + "))"
                else:
                    raise HTTPException(status_code=422, detail="source_wkt does not match geometry_kind")
            elif kind != "area":
                longitude, latitude = parse_coordinates(
                    f"{(row.get('longitude') or '').strip()},{(row.get('latitude') or '').strip()}"
                )
                wkt = f"POINT({longitude} {latitude})"
            elif kind == "area":
                vertices = parse_polygon_vertices(row.get("polygon_vertices") or "")
                wkt = "POLYGON((" + ", ".join(f"{lon} {lat}" for lon, lat in vertices) + "))"
        except HTTPException as exc:
            errors.append(exc.detail)
        inventory = (row.get("inventory_number") or "").strip() or None
        try:
            year = int(row["year"]) if (row.get("year") or "").strip() else None
            if year is not None and not 1500 <= year <= 2200:
                raise ValueError
            start_year = int(row["work_year_start"]) if (row.get("work_year_start") or "").strip() else None
            end_year = int(row["work_year_end"]) if (row.get("work_year_end") or "").strip() else None
            if any(value is not None and not 1500 <= value <= 2200 for value in (start_year, end_year)):
                raise ValueError
            if start_year and end_year and start_year > end_year:
                raise ValueError
        except ValueError:
            errors.append("year/work_year_start/work_year_end must be valid years (1500–2200)")
            year = start_year = end_year = None
        efgi_url = (row.get("efgi_url") or "").strip() or None
        if efgi_url and not is_http_url(efgi_url):
            errors.append("efgi_url must use HTTP or HTTPS")
        length_limits = {
            "inventory_number": 200, "tgf_number": 200, "region": 200, "topic": 1000, "description": 10000,
            "archive_reference": 1000, "source_wkt": 20000, "document_type": 200,
            "authors": 2000, "coauthors": 2000,
            "executor_org": 500, "created_place": 500, "minerals": 2000, "archive_disk_number": 200,
            "material_composition": 4000, "electronic_copy_status": 200, "efgi_id": 200, "efgi_url": 2000,
        }
        for field, max_length in length_limits.items():
            if row.get(field) and len(row[field]) > max_length:
                errors.append(f"{field} exceeds {max_length} characters")
        validated.append({
            "row": index,
            "feature_name": feature_name,
            "geometry_kind": kind,
            "coordinate_crs": crs,
            "wkt": wkt,
            "document_title": title,
            "inventory_number": inventory,
            "tgf_number": (row.get("tgf_number") or "").strip() or None,
            "region": (row.get("region") or "").strip() or None,
            "year": year,
            "topic": (row.get("topic") or "").strip() or None,
            "description": (row.get("description") or "").strip() or None,
            "archive_reference": (row.get("archive_reference") or "").strip() or None,
            "document_type": (row.get("document_type") or "").strip() or None,
            "authors": (row.get("authors") or "").strip() or None,
            "coauthors": (row.get("coauthors") or "").strip() or None,
            "executor_org": (row.get("executor_org") or "").strip() or None,
            "work_year_start": start_year,
            "work_year_end": end_year,
            "created_place": (row.get("created_place") or "").strip() or None,
            "minerals": (row.get("minerals") or "").strip() or None,
            "archive_disk_number": (row.get("archive_disk_number") or "").strip() or None,
            "material_composition": (row.get("material_composition") or "").strip() or None,
            "electronic_copy_status": (row.get("electronic_copy_status") or "").strip() or None,
            "efgi_id": (row.get("efgi_id") or "").strip() or None,
            "efgi_url": efgi_url,
            "errors": errors,
        })
    for field, label in (("inventory_number", "inventory number"),):
        values = [row[field] for row in validated if row[field]]
        duplicates = {value for value in values if values.count(value) > 1}
        for row in validated:
            if row[field] in duplicates:
                row["errors"].append(f"duplicate {label} within CSV")
    return validated


def parse_catalog_xlsx(raw: bytes) -> list[dict]:
    if len(raw) > max_csv_import_size:
        raise HTTPException(status_code=413, detail="XLSX import cannot exceed 5 MiB")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if (
                len(members) > 10000
                or sum(member.file_size for member in members) > 25 * 1024 * 1024
                or any(member.filename.startswith("/") or ".." in Path(member.filename).parts for member in members)
                or any(member.filename.lower().endswith("vbaproject.bin") for member in members)
            ):
                raise ValueError("workbook contents exceed import limits or use an unsupported format")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.active
            if sheet.max_row is not None and sheet.max_row > max_csv_import_rows + 1:
                raise ValueError(f"worksheet cannot exceed {max_csv_import_rows} data rows")
            if sheet.max_column is not None and sheet.max_column > 40:
                raise ValueError("worksheet cannot exceed 40 columns")
            csv_buffer = io.StringIO(newline="")
            writer = csv.writer(csv_buffer)
            for row_number, cells in enumerate(
                sheet.iter_rows(max_row=max_csv_import_rows + 2, max_col=sheet.max_column or 1), 1
            ):
                if row_number > max_csv_import_rows + 1:
                    raise ValueError(f"worksheet cannot exceed {max_csv_import_rows} data rows")
                values = []
                for cell in cells:
                    if cell.data_type == "f":
                        raise ValueError("formulas are not allowed in import workbooks")
                    value = cell.value
                    if isinstance(value, (datetime, date)):
                        value = value.isoformat()
                    elif value is not None and not isinstance(value, str):
                        value = str(value)
                    values.append("" if value is None else value)
                if any(str(value).strip() for value in values):
                    writer.writerow(values)
            return parse_catalog_csv(csv_buffer.getvalue().encode("utf-8"))
        finally:
            workbook.close()
    except HTTPException:
        raise
    except (zipfile.BadZipFile, InvalidFileException, OSError, ValueError, KeyError, IndexError, ParseError) as exc:
        raise HTTPException(status_code=422, detail=f"invalid XLSX workbook: {exc}") from exc


def parse_catalog_import(raw: bytes, file_format: str) -> list[dict]:
    if file_format == "csv":
        return parse_catalog_csv(raw)
    if file_format == "xlsx":
        return parse_catalog_xlsx(raw)
    raise HTTPException(status_code=415, detail="import format must be csv or xlsx")


@app.get("/")
def index():
    return FileResponse("app/static/index.html")


@app.get("/login")
def login_page():
    return FileResponse("app/static/login.html")


@app.get("/admin/users")
def users_page(request: Request):
    require_admin(request)
    return FileResponse("app/static/users.html")


@app.get("/admin/catalog")
def catalog_admin_page(request: Request):
    require_admin(request)
    return FileResponse("app/static/catalog-admin.html")


@app.get("/api/session")
def get_session(request: Request):
    user = getattr(request.state, "user", None)
    if user is None:
        raise HTTPException(status_code=401, detail="authentication required")
    return {"id": user["id"], "username": user["username"], "role": user["role"]}


@app.get("/api/saved-searches")
def list_saved_searches(request: Request):
    user = request.state.user
    with closing(get_connection()) as connection:
        searches = connection.execute(
            """
            SELECT id, name, filters, created_at
            FROM saved_searches
            WHERE user_id = %s
            ORDER BY name
            """,
            (user["id"],),
        ).fetchall()
    return {"searches": searches}


@app.post("/api/saved-searches", status_code=201)
def create_saved_search(search: SavedSearchCreate, request: Request):
    require_same_origin(request)
    filters = {key: value for key, value in search.filters.items() if str(value).strip()}
    try:
        with get_connection() as connection:
            created = connection.execute(
                """
                INSERT INTO saved_searches (user_id, name, filters)
                VALUES (%s, %s, %s)
                RETURNING id, name, filters, created_at
                """,
                (request.state.user["id"], search.name.strip(), json.dumps(filters)),
            ).fetchone()
    except UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="a saved search with this name already exists") from exc
    return created


@app.delete("/api/saved-searches/{search_id}", status_code=204)
def delete_saved_search(search_id: int, request: Request):
    require_same_origin(request)
    with get_connection() as connection:
        deleted = connection.execute(
            "DELETE FROM saved_searches WHERE id = %s AND user_id = %s RETURNING id",
            (search_id, request.state.user["id"]),
        ).fetchone()
    if deleted is None:
        raise HTTPException(status_code=404, detail="saved search not found")
    return None


@app.patch("/api/admin/documents/{document_id}")
def update_document(document_id: int, update: DocumentUpdate, request: Request):
    user = require_admin(request)
    require_same_origin(request)
    changes = update.model_dump(exclude_unset=True)
    for key, value in changes.items():
        if isinstance(value, str):
            changes[key] = value.strip() or None
    if "title" in changes and changes["title"] is None:
        raise HTTPException(status_code=422, detail="document title cannot be blank")
    if changes.get("efgi_url") and not is_http_url(changes["efgi_url"]):
        raise HTTPException(status_code=422, detail="efgi_url must use HTTP or HTTPS")
    assignments = [f"{key} = %s" for key in changes]
    values = list(changes.values()) + [document_id]
    try:
        with get_connection() as connection:
            current = connection.execute(
                "SELECT work_year_start, work_year_end FROM documents WHERE id = %s",
                (document_id,),
            ).fetchone()
            if current is None:
                raise HTTPException(status_code=404, detail="document not found")
            start_year = changes.get("work_year_start", current["work_year_start"])
            end_year = changes.get("work_year_end", current["work_year_end"])
            if start_year and end_year and start_year > end_year:
                raise HTTPException(status_code=422, detail="work start year must not exceed end year")
            updated = connection.execute(
                f"""
                UPDATE documents SET {", ".join(assignments)}
                WHERE id = %s
                RETURNING id, title, inventory_number, region, year, topic, description, archive_reference,
                          tgf_number, document_type, authors, coauthors, executor_org, work_year_start, work_year_end,
                          created_place, minerals, archive_disk_number, material_composition,
                          electronic_copy_status, efgi_id, efgi_url
                """,
                values,
            ).fetchone()
            if updated is None:
                raise HTTPException(status_code=404, detail="document not found")
            write_audit(connection, user["id"], "update", "document", document_id, {"fields": sorted(changes)})
    except UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="inventory number already exists") from exc
    return updated


@app.post("/api/admin/documents/{document_id}/relations", status_code=201)
def create_document_relation(document_id: int, relation: DocumentRelationCreate, request: Request):
    user = require_admin(request)
    require_same_origin(request)
    if document_id == relation.target_document_id:
        raise HTTPException(status_code=422, detail="a document cannot be related to itself")
    try:
        with get_connection() as connection:
            created = connection.execute(
                """
                INSERT INTO document_relations (
                    source_document_id, target_document_id, relation_type, created_by
                )
                VALUES (%s, %s, %s, %s)
                RETURNING id, source_document_id, target_document_id, relation_type, created_at
                """,
                (document_id, relation.target_document_id, relation.relation_type, user["id"]),
            ).fetchone()
            if created is None:
                raise HTTPException(status_code=404, detail="document not found")
            write_audit(
                connection, user["id"], "link", "document_relation", created["id"],
                {"source_document_id": document_id, "target_document_id": relation.target_document_id,
                 "relation_type": relation.relation_type},
            )
    except psycopg.errors.ForeignKeyViolation as exc:
        raise HTTPException(status_code=404, detail="one of the documents was not found") from exc
    except UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="this document relationship already exists") from exc
    return created


@app.delete("/api/admin/document-relations/{relation_id}", status_code=204)
def delete_document_relation(relation_id: int, request: Request):
    user = require_admin(request)
    require_same_origin(request)
    with get_connection() as connection:
        relation = connection.execute(
            "DELETE FROM document_relations WHERE id = %s RETURNING source_document_id, target_document_id",
            (relation_id,),
        ).fetchone()
        if relation is None:
            raise HTTPException(status_code=404, detail="document relationship not found")
        write_audit(connection, user["id"], "unlink", "document_relation", relation_id, relation)
    return None


@app.get("/api/admin/audit")
def list_catalog_audit(request: Request, limit: Annotated[int, Query(ge=1, le=500)] = 100):
    require_admin(request)
    with closing(get_connection()) as connection:
        entries = connection.execute(
            """
            SELECT audit.id, audit.action, audit.entity_type, audit.entity_id, audit.details,
                   audit.created_at, users.username AS actor
            FROM catalog_audit audit
            LEFT JOIN geocatalog_users users ON users.id = audit.actor_id
            ORDER BY audit.created_at DESC, audit.id DESC
            LIMIT %s
            """,
            (limit,),
        ).fetchall()
    return {"entries": entries}


@app.get("/api/admin/quality")
def catalog_quality_report(request: Request):
    require_admin(request)
    with closing(get_connection()) as connection:
        unmapped = connection.execute(
            """
            SELECT id, name, source_crs FROM features
            WHERE geom IS NULL ORDER BY name LIMIT 500
            """
        ).fetchall()
        unlinked_documents = connection.execute(
            """
            SELECT d.id, d.title, d.inventory_number FROM documents d
            LEFT JOIN feature_documents fd ON fd.document_id = d.id
            WHERE fd.document_id IS NULL ORDER BY d.title LIMIT 500
            """
        ).fetchall()
        incomplete_documents = connection.execute(
            """
            SELECT id, title, inventory_number FROM documents
            WHERE title IS NULL OR btrim(title) = '' OR
                  (inventory_number IS NULL AND tgf_number IS NULL AND document_type IS NULL AND year IS NULL)
            ORDER BY title LIMIT 500
            """
        ).fetchall()
        duplicates = connection.execute(
            """
            SELECT inventory_number, count(*) AS count
            FROM documents WHERE inventory_number IS NOT NULL
            GROUP BY inventory_number HAVING count(*) > 1
            ORDER BY inventory_number LIMIT 500
            """
        ).fetchall()
        duplicate_tgf_numbers = connection.execute(
            """
            SELECT tgf_number, count(*) AS count
            FROM documents WHERE tgf_number IS NOT NULL
            GROUP BY tgf_number HAVING count(*) > 1
            ORDER BY tgf_number LIMIT 500
            """
        ).fetchall()
        totals = connection.execute(
            """
            SELECT
                (SELECT count(*) FROM features WHERE geom IS NULL) AS unmapped_features,
                (SELECT count(*) FROM documents d LEFT JOIN feature_documents fd ON fd.document_id = d.id
                 WHERE fd.document_id IS NULL) AS unlinked_documents,
                (SELECT count(*) FROM documents
                 WHERE title IS NULL OR btrim(title) = '' OR
                       (inventory_number IS NULL AND tgf_number IS NULL AND document_type IS NULL AND year IS NULL))
                    AS incomplete_documents,
                (SELECT count(*) FROM (
                    SELECT inventory_number FROM documents WHERE inventory_number IS NOT NULL
                    GROUP BY inventory_number HAVING count(*) > 1
                ) duplicate_inventory) AS duplicate_inventory_numbers,
                (SELECT count(*) FROM (
                    SELECT tgf_number FROM documents WHERE tgf_number IS NOT NULL
                    GROUP BY tgf_number HAVING count(*) > 1
                ) duplicate_tgf) AS duplicate_tgf_numbers
            """
        ).fetchone()
    return {
        "unmapped_features": unmapped,
        "unlinked_documents": unlinked_documents,
        "incomplete_documents": incomplete_documents,
        "duplicate_inventory_numbers": duplicates,
        "duplicate_tgf_numbers": duplicate_tgf_numbers,
        "counts": {
            key: totals[key] for key in totals
        },
    }


@app.get("/api/export.csv")
def export_catalog_csv(
    request: Request,
    q: Annotated[str | None, Query(max_length=200)] = None,
    region: Annotated[str | None, Query(max_length=200)] = None,
    year_from: Annotated[int | None, Query(ge=1500, le=2200)] = None,
    year_to: Annotated[int | None, Query(ge=1500, le=2200)] = None,
):
    if year_from is not None and year_to is not None and year_from > year_to:
        raise HTTPException(status_code=422, detail="year_from must not exceed year_to")
    conditions = []
    params = []
    if q:
        conditions.append(
            "(d.search_vector @@ websearch_to_tsquery('simple', %s) "
            "OR d.search_vector @@ websearch_to_tsquery('russian', %s) "
            "OR d.search_vector @@ websearch_to_tsquery('english', %s) "
            "OR d.authors ILIKE %s OR d.coauthors ILIKE %s OR d.executor_org ILIKE %s "
            "OR d.document_type ILIKE %s OR d.minerals ILIKE %s OR d.archive_disk_number ILIKE %s "
            "OR d.tgf_number ILIKE %s)"
        )
        params.extend((q, q, q, *(f"%{q}%" for _ in range(7))))
    if region:
        conditions.append("d.region ILIKE %s")
        params.append(f"%{region}%")
    if year_from is not None:
        conditions.append("coalesce(d.work_year_start, d.year) >= %s")
        params.append(year_from)
    if year_to is not None:
        conditions.append("coalesce(d.work_year_end, d.year) <= %s")
        params.append(year_to)
    where_clause = " AND ".join(conditions) if conditions else "TRUE"
    with closing(get_connection()) as connection:
        rows = connection.execute(
            f"""
            SELECT f.name AS feature_name, f.kind, f.source_crs,
                   ST_AsText(f.source_geom) AS source_wkt,
                   d.title, d.inventory_number, d.region, d.year, d.topic, d.description,
                   d.archive_reference, d.tgf_number, d.document_type, d.authors, d.coauthors, d.executor_org,
                   d.work_year_start, d.work_year_end, d.created_place, d.minerals,
                   d.archive_disk_number, d.material_composition, d.electronic_copy_status,
                   d.efgi_id, d.efgi_url
            FROM documents d
            JOIN feature_documents fd ON fd.document_id = d.id
            JOIN features f ON f.id = fd.feature_id
            WHERE {where_clause}
            ORDER BY f.name, d.year DESC NULLS LAST, d.title
            """,
            params,
        ).fetchall()
    columns = [
        "feature_name", "geometry_kind", "coordinate_crs", "source_wkt", "document_title",
        "inventory_number", "tgf_number", "region", "year", "topic", "description", "archive_reference",
        "document_type", "authors", "coauthors", "executor_org", "work_year_start", "work_year_end",
        "created_place", "minerals", "archive_disk_number", "material_composition",
        "electronic_copy_status", "efgi_id", "efgi_url",
    ]
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([
            csv_export_value(value)
            for value in [
                row["feature_name"], row["kind"], row["source_crs"], row["source_wkt"],
                row["title"], row["inventory_number"], row["tgf_number"], row["region"], row["year"],
                row["topic"], row["description"], row["archive_reference"], row["document_type"], row["authors"],
                row["coauthors"], row["executor_org"], row["work_year_start"], row["work_year_end"],
                row["created_place"], row["minerals"], row["archive_disk_number"],
                row["material_composition"], row["electronic_copy_status"], row["efgi_id"], row["efgi_url"],
            ]
        ])
    return Response(
        "\ufeff" + output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="geocatalog-export.csv"'},
    )


@app.post("/api/admin/import/{file_format}/preview")
async def preview_catalog_import(file_format: str, request: Request):
    require_admin(request)
    require_same_origin(request)
    raw = await request.body()
    rows = parse_catalog_import(raw, file_format)
    return {"rows": rows, "valid": all(not row["errors"] for row in rows), "count": len(rows)}


@app.post("/api/admin/import/{file_format}/commit", status_code=201)
async def commit_catalog_import(file_format: str, request: Request):
    user = require_admin(request)
    require_same_origin(request)
    rows = parse_catalog_import(await request.body(), file_format)
    invalid_rows = [row["row"] for row in rows if row["errors"]]
    if invalid_rows:
        raise HTTPException(status_code=422, detail={"message": "CSV has validation errors", "rows": invalid_rows})
    inserted_documents = []
    skipped_documents = 0
    try:
        with get_connection() as connection:
            features_by_name = {}
            for row in rows:
                fingerprint = import_fingerprint(row)
                existing_import = connection.execute(
                    "SELECT id FROM documents WHERE import_fingerprint = %s",
                    (fingerprint,),
                ).fetchone()
                if existing_import is not None:
                    skipped_documents += 1
                    continue
                feature = features_by_name.get(row["feature_name"])
                if feature is None:
                    feature = connection.execute(
                        """
                        SELECT id, name, kind, source_crs, ST_AsText(source_geom) AS coordinates
                        FROM features WHERE name = %s
                        """,
                        (row["feature_name"],),
                    ).fetchone()
                    if feature and (
                        feature["kind"] != row["geometry_kind"]
                        or feature["source_crs"] != row["coordinate_crs"]
                        or feature["coordinates"] != row["wkt"]
                    ):
                        raise HTTPException(
                            status_code=409,
                            detail=f"row {row['row']}: feature name already exists with different geometry",
                        )
                    if feature is None:
                        crs_srid = 7683 if row["coordinate_crs"] == "EPSG:7683" else 4326
                        valid_geometry = connection.execute(
                            "SELECT ST_IsValid(ST_GeomFromText(%s, %s)) AS valid",
                            (row["wkt"], crs_srid),
                        ).fetchone()["valid"]
                        if not valid_geometry:
                            raise HTTPException(
                                status_code=422,
                                detail=f"row {row['row']}: area contour is self-intersecting or invalid",
                            )
                        feature = connection.execute(
                            """
                            INSERT INTO features (name, kind, geom, metadata, source_geom, source_crs)
                            VALUES (
                                %s, %s,
                                CASE WHEN %s = 4326 THEN ST_GeomFromText(%s, 4326) ELSE NULL END,
                                '{}'::jsonb, ST_GeomFromText(%s, %s), %s
                            )
                            RETURNING id, name, kind, source_crs
                            """,
                            (
                                row["feature_name"], row["geometry_kind"], crs_srid, row["wkt"],
                                row["wkt"], crs_srid, row["coordinate_crs"],
                            ),
                        ).fetchone()
                    features_by_name[row["feature_name"]] = feature
                document = connection.execute(
                    """
                    INSERT INTO documents (
                        title, inventory_number, region, year, topic, description, archive_reference,
                        tgf_number, document_type, authors, coauthors, executor_org, work_year_start, work_year_end,
                        created_place, minerals, archive_disk_number, material_composition,
                        electronic_copy_status, efgi_id, efgi_url, import_fingerprint
                    )
                    VALUES (
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
                    )
                    RETURNING id
                    """,
                    (
                        row["document_title"], row["inventory_number"], row["region"], row["year"],
                        row["topic"], row["description"], row["archive_reference"], row["tgf_number"],
                        row["document_type"], row["authors"], row["coauthors"], row["executor_org"], row["work_year_start"],
                        row["work_year_end"], row["created_place"], row["minerals"],
                        row["archive_disk_number"], row["material_composition"],
                        row["electronic_copy_status"], row["efgi_id"], row["efgi_url"], fingerprint,
                    ),
                ).fetchone()
                connection.execute(
                    "INSERT INTO feature_documents (feature_id, document_id) VALUES (%s, %s)",
                    (feature["id"], document["id"]),
                )
                inserted_documents.append(document["id"])
            write_audit(
                connection, user["id"], "import", "documents", None,
                {"count": len(inserted_documents), "skipped_repeats": skipped_documents, "format": file_format},
            )
    except UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="duplicate feature name or inventory number; import rolled back") from exc
    return {
        "imported_documents": len(inserted_documents),
        "skipped_repeats": skipped_documents,
        "document_ids": inserted_documents,
    }


@app.get("/api/admin/users")
def list_users(request: Request):
    require_admin(request)
    with closing(get_connection()) as connection:
        users = connection.execute(
            """
            SELECT id, username, role, is_active, created_at
            FROM geocatalog_users
            ORDER BY username
            """
        ).fetchall()
    return {"users": users}


@app.post("/api/admin/users", status_code=201)
def create_user(user: UserCreate, request: Request):
    require_admin(request)
    try:
        with get_connection() as connection:
            created = connection.execute(
                """
                INSERT INTO geocatalog_users (username, password_hash, role)
                VALUES (%s, %s, %s)
                RETURNING id, username, role, is_active, created_at
                """,
                (user.username, password_hasher.hash(user.password), user.role),
            ).fetchone()
    except UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="username already exists") from exc
    return created


@app.patch("/api/admin/users/{user_id}")
def update_user(user_id: int, update: UserUpdate, request: Request):
    require_admin(request)
    changes = update.model_dump(exclude_unset=True)
    with get_connection() as connection:
        connection.execute("SELECT pg_advisory_xact_lock(73412905)")
        current = connection.execute(
            "SELECT id, username, role, is_active FROM geocatalog_users WHERE id = %s FOR UPDATE",
            (user_id,),
        ).fetchone()
        if current is None:
            raise HTTPException(status_code=404, detail="user not found")

        next_role = changes.get("role", current["role"])
        next_active = changes.get("active", current["is_active"])
        if current["role"] == "admin" and current["is_active"] and (
            next_role != "admin" or not next_active
        ):
            admin_count = connection.execute(
                "SELECT count(*) AS count FROM geocatalog_users WHERE role = 'admin' AND is_active"
            ).fetchone()["count"]
            if admin_count <= 1:
                raise HTTPException(status_code=409, detail="cannot disable or demote the last active administrator")

        assignments = []
        values = []
        if "role" in changes:
            assignments.append("role = %s")
            values.append(changes["role"])
        if "active" in changes:
            assignments.append("is_active = %s")
            values.append(changes["active"])
        if "password" in changes:
            assignments.extend(("password_hash = %s", "token_version = token_version + 1"))
            values.append(password_hasher.hash(changes["password"]))
        values.append(user_id)
        updated = connection.execute(
            f"""
            UPDATE geocatalog_users
            SET {", ".join(assignments)}
            WHERE id = %s
            RETURNING id, username, role, is_active, created_at
            """,
            values,
        ).fetchone()
    return updated


@app.get("/api/admin/features")
def list_features_for_admin(request: Request):
    require_admin(request)
    with closing(get_connection()) as connection:
        features = connection.execute(
            "SELECT id, name, kind FROM features ORDER BY name"
        ).fetchall()
    return {"features": features}


@app.get("/api/admin/documents")
def list_documents_for_admin(request: Request):
    require_admin(request)
    with closing(get_connection()) as connection:
        documents = connection.execute(
            """
            SELECT id, title, inventory_number FROM documents
            ORDER BY title, inventory_number NULLS LAST LIMIT 1000
            """
        ).fetchall()
    return {"documents": documents}


@app.get("/api/admin/documents/{document_id}")
def get_document_for_admin(document_id: int, request: Request):
    require_admin(request)
    with closing(get_connection()) as connection:
        document = connection.execute(
            """
            SELECT id, title, inventory_number, tgf_number, region, year, topic, description,
                   archive_reference, document_type, authors, coauthors, executor_org,
                   work_year_start, work_year_end, created_place, minerals, archive_disk_number,
                   material_composition, electronic_copy_status, efgi_id, efgi_url
            FROM documents WHERE id = %s
            """,
            (document_id,),
        ).fetchone()
    if document is None:
        raise HTTPException(status_code=404, detail="document not found")
    return document


@app.get("/api/admin/files")
def list_files_for_admin(request: Request):
    require_admin(request)
    with closing(get_connection()) as connection:
        files = connection.execute(
            """
            SELECT file.id, file.original_filename, file.size_bytes, document.title AS document_title,
                   feature.name AS feature_name
            FROM document_files file
            JOIN documents document ON document.id = file.document_id
            JOIN feature_documents link ON link.document_id = document.id
            JOIN features feature ON feature.id = link.feature_id
            ORDER BY file.created_at DESC, file.original_filename
            """
        ).fetchall()
    return {"files": files}


@app.delete("/api/admin/files/{file_id}", status_code=204)
def delete_uploaded_file(file_id: int, request: Request):
    user = require_admin(request)
    require_same_origin(request)
    with get_connection() as connection:
        file = connection.execute(
            "DELETE FROM document_files WHERE id = %s RETURNING storage_key",
            (file_id,),
        ).fetchone()
        if file is not None:
            write_audit(connection, user["id"], "delete", "document_file", file_id, {})
    if file is None:
        raise HTTPException(status_code=404, detail="file not found")
    if not re.fullmatch(r"[a-f0-9]{48}\.[a-z0-9]+", file["storage_key"]):
        raise HTTPException(status_code=500, detail="invalid stored file key")
    (upload_root / file["storage_key"]).unlink(missing_ok=True)
    return None


@app.post("/api/admin/documents", status_code=201)
async def create_document_with_files(
    request: Request,
    feature_mode: Annotated[str, Form()],
    document_title: Annotated[str, Form(min_length=1, max_length=500)],
    files: Annotated[list[UploadFile], File()],
    feature_id: Annotated[str, Form()] = "",
    feature_name: Annotated[str, Form(max_length=500)] = "",
    geometry_kind: Annotated[str, Form()] = "well",
    longitude: Annotated[str, Form()] = "",
    latitude: Annotated[str, Form()] = "",
    coordinate_crs: Annotated[str, Form()] = "EPSG:7683",
    polygon_vertices: Annotated[str, Form(max_length=20000)] = "",
    feature_metadata: Annotated[str, Form(max_length=10000)] = "{}",
    inventory_number: Annotated[str, Form(max_length=200)] = "",
    region: Annotated[str, Form(max_length=200)] = "",
    document_year: Annotated[str, Form()] = "",
    topic: Annotated[str, Form(max_length=1000)] = "",
    description: Annotated[str, Form(max_length=10000)] = "",
    archive_reference: Annotated[str, Form(max_length=1000)] = "",
    tgf_number: Annotated[str, Form(max_length=200)] = "",
    document_type: Annotated[str, Form(max_length=200)] = "",
    authors: Annotated[str, Form(max_length=2000)] = "",
    coauthors: Annotated[str, Form(max_length=2000)] = "",
    executor_org: Annotated[str, Form(max_length=500)] = "",
    work_year_start: Annotated[str, Form()] = "",
    work_year_end: Annotated[str, Form()] = "",
    created_place: Annotated[str, Form(max_length=500)] = "",
    minerals: Annotated[str, Form(max_length=2000)] = "",
    archive_disk_number: Annotated[str, Form(max_length=200)] = "",
    material_composition: Annotated[str, Form(max_length=4000)] = "",
    electronic_copy_status: Annotated[str, Form(max_length=200)] = "",
    efgi_id: Annotated[str, Form(max_length=200)] = "",
    efgi_url: Annotated[str, Form(max_length=2000)] = "",
    related_document_id: Annotated[str, Form()] = "",
    relation_type: Annotated[str, Form()] = "related",
):
    require_admin(request)
    require_same_origin(request)
    if feature_mode not in {"new", "existing"}:
        raise HTTPException(status_code=422, detail="feature_mode must be new or existing")
    title = document_title.strip()
    if not title:
        raise HTTPException(status_code=422, detail="document title cannot be blank")
    if not files or len(files) > max_upload_files:
        raise HTTPException(status_code=413, detail="select between 1 and 10 files")

    year = None
    if document_year.strip():
        try:
            year = int(document_year)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="document year must be a number") from exc
        if not 1500 <= year <= 2200:
            raise HTTPException(status_code=422, detail="document year must be between 1500 and 2200")
    work_years = []
    for value in (work_year_start, work_year_end):
        if value.strip():
            try:
                parsed_year = int(value)
            except ValueError as exc:
                raise HTTPException(status_code=422, detail="work years must be numbers") from exc
            if not 1500 <= parsed_year <= 2200:
                raise HTTPException(status_code=422, detail="work years must be between 1500 and 2200")
            work_years.append(parsed_year)
        else:
            work_years.append(None)
    if work_years[0] and work_years[1] and work_years[0] > work_years[1]:
        raise HTTPException(status_code=422, detail="work start year must not exceed end year")
    if efgi_url.strip() and not is_http_url(efgi_url.strip()):
        raise HTTPException(status_code=422, detail="EFGI URL must use HTTP or HTTPS")
    related_id = None
    if related_document_id.strip():
        try:
            related_id = int(related_document_id)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="related document must be a valid document ID") from exc
        if relation_type not in {"related", "appendix", "protocol", "map", "copy_of"}:
            raise HTTPException(status_code=422, detail="invalid document relationship type")

    prepared_files = []
    request_size = 0
    for upload in files:
        data = await upload.read(max_file_size + 1)
        filename, media_type, extension = validate_uploaded_file(upload.filename, data)
        request_size += len(data)
        if request_size > max_request_size:
            raise HTTPException(status_code=413, detail="total file size per upload cannot exceed 100 MiB")
        storage_key = f"{secrets.token_hex(24)}.{extension}"
        prepared_files.append((filename, media_type, storage_key, data))

    if feature_mode == "existing":
        try:
            selected_feature_id = int(feature_id)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="select an existing catalog object") from exc
        new_feature_name = None
        geometry_wkt = None
        metadata = None
    else:
        new_feature_name = feature_name.strip()
        if not new_feature_name:
            raise HTTPException(status_code=422, detail="new catalog object needs a name")
        if coordinate_crs not in {"EPSG:7683", "EPSG:4326"}:
            raise HTTPException(status_code=422, detail="coordinate CRS must be EPSG:7683 or EPSG:4326")
        if geometry_kind == "well":
            try:
                longitude_value, latitude_value = parse_coordinates(f"{longitude},{latitude}")
            except HTTPException:
                raise
            geometry_wkt = None
        elif geometry_kind == "area":
            vertices = parse_polygon_vertices(polygon_vertices)
            geometry_wkt = "POLYGON((" + ", ".join(f"{lon} {lat}" for lon, lat in vertices) + "))"
            longitude_value = latitude_value = None
        else:
            raise HTTPException(status_code=422, detail="geometry kind must be a point or area")
        try:
            metadata = json.loads(feature_metadata or "{}")
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=422, detail="feature metadata must be valid JSON") from exc
        if not isinstance(metadata, dict):
            raise HTTPException(status_code=422, detail="feature metadata must be a JSON object")
        selected_feature_id = None

    upload_root.mkdir(parents=True, exist_ok=True)
    saved_paths = []
    try:
        for _, _, storage_key, data in prepared_files:
            path = upload_root / storage_key
            descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            saved_paths.append(path)
            with os.fdopen(descriptor, "wb") as stored_file:
                stored_file.write(data)

        try:
            with get_connection() as connection:
                if feature_mode == "existing":
                    feature = connection.execute(
                        "SELECT id, name FROM features WHERE id = %s",
                        (selected_feature_id,),
                    ).fetchone()
                    if feature is None:
                        raise HTTPException(status_code=404, detail="catalog object not found")
                elif geometry_kind == "well":
                    source_srid = 7683 if coordinate_crs == "EPSG:7683" else 4326
                    feature = connection.execute(
                        """
                        INSERT INTO features (name, kind, geom, metadata, source_geom, source_crs)
                        VALUES (
                            %s, 'well',
                            CASE WHEN %s = 4326 THEN ST_SetSRID(ST_MakePoint(%s, %s), 4326) ELSE NULL END,
                            %s, ST_SetSRID(ST_MakePoint(%s, %s), %s), %s
                        )
                        RETURNING id, name
                        """,
                        (
                            new_feature_name, source_srid, longitude_value, latitude_value,
                            json.dumps(metadata), longitude_value, latitude_value, source_srid, coordinate_crs,
                        ),
                    ).fetchone()
                else:
                    source_srid = 7683 if coordinate_crs == "EPSG:7683" else 4326
                    valid = connection.execute(
                        "SELECT ST_IsValid(ST_GeomFromText(%s, %s)) AS valid",
                        (geometry_wkt, source_srid),
                    ).fetchone()["valid"]
                    if not valid:
                        raise HTTPException(status_code=422, detail="area contour is self-intersecting or invalid")
                    feature = connection.execute(
                        """
                        INSERT INTO features (name, kind, geom, metadata, source_geom, source_crs)
                        VALUES (
                            %s, 'area',
                            CASE WHEN %s = 4326 THEN ST_GeomFromText(%s, 4326) ELSE NULL END,
                            %s, ST_GeomFromText(%s, %s), %s
                        )
                        RETURNING id, name
                        """,
                        (
                            new_feature_name, source_srid, geometry_wkt, json.dumps(metadata),
                            geometry_wkt, source_srid, coordinate_crs,
                        ),
                    ).fetchone()

                document = connection.execute(
                    """
                    INSERT INTO documents (
                        title, inventory_number, region, year, topic, description, archive_reference,
                        tgf_number, document_type, authors, coauthors, executor_org, work_year_start, work_year_end,
                        created_place, minerals, archive_disk_number, material_composition,
                        electronic_copy_status, efgi_id, efgi_url
                    )
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    RETURNING id
                    """,
                    (
                        title,
                        inventory_number.strip() or None,
                        region.strip() or None,
                        year,
                        topic.strip() or None,
                        description.strip() or None,
                        archive_reference.strip() or None,
                        tgf_number.strip() or None,
                        document_type.strip() or None,
                        authors.strip() or None,
                        coauthors.strip() or None,
                        executor_org.strip() or None,
                        work_years[0],
                        work_years[1],
                        created_place.strip() or None,
                        minerals.strip() or None,
                        archive_disk_number.strip() or None,
                        material_composition.strip() or None,
                        electronic_copy_status.strip() or None,
                        efgi_id.strip() or None,
                        efgi_url.strip() or None,
                    ),
                ).fetchone()
                connection.execute(
                    "INSERT INTO feature_documents (feature_id, document_id) VALUES (%s, %s)",
                    (feature["id"], document["id"]),
                )
                if related_id is not None:
                    if related_id == document["id"]:
                        raise HTTPException(status_code=422, detail="a document cannot be related to itself")
                    connection.execute(
                        """
                        INSERT INTO document_relations (
                            source_document_id, target_document_id, relation_type, created_by
                        )
                        VALUES (%s, %s, %s, %s)
                        """,
                        (document["id"], related_id, relation_type, request.state.user["id"]),
                    )
                for filename, media_type, storage_key, data in prepared_files:
                    connection.execute(
                        """
                        INSERT INTO document_files (document_id, storage_key, original_filename, media_type, size_bytes)
                        VALUES (%s, %s, %s, %s, %s)
                        """,
                        (document["id"], storage_key, filename, media_type, len(data)),
                    )
                write_audit(
                    connection, request.state.user["id"], "create", "document", document["id"],
                    {
                        "feature_id": feature["id"], "file_count": len(prepared_files),
                        "coordinate_crs": coordinate_crs if feature_mode == "new" else None,
                        "related_document_id": related_id,
                    },
                )
        except psycopg.errors.ForeignKeyViolation as exc:
            raise HTTPException(status_code=404, detail="related document was not found") from exc
        except UniqueViolation as exc:
            raise HTTPException(status_code=409, detail="object name or inventory number already exists") from exc
    except Exception:
        for path in saved_paths:
            path.unlink(missing_ok=True)
        raise

    return {
        "feature_id": feature["id"],
        "feature_name": feature["name"],
        "document_id": document["id"],
        "files": [{"filename": filename, "size_bytes": len(data)} for filename, _, _, data in prepared_files],
    }


@app.get("/api/files/{file_id}")
def download_document_file(file_id: int):
    with closing(get_connection()) as connection:
        file = connection.execute(
            """
            SELECT storage_key, original_filename
            FROM document_files
            WHERE id = %s
            """,
            (file_id,),
        ).fetchone()
    if file is None or not re.fullmatch(r"[a-f0-9]{48}\.[a-z0-9]+", file["storage_key"]):
        raise HTTPException(status_code=404, detail="file not found")
    path = upload_root / file["storage_key"]
    if not path.is_file():
        raise HTTPException(status_code=404, detail="stored file is missing")
    return FileResponse(
        path,
        media_type="application/octet-stream",
        filename=file["original_filename"],
        headers={"Cache-Control": "private, no-store"},
    )


@app.post("/auth/login")
async def login(request: Request):
    if request.url.scheme != "https":
        raise HTTPException(status_code=426, detail="HTTPS is required for login")
    host = request.headers.get("host", "")
    origin = request.headers.get("origin")
    if origin and urlsplit(origin).netloc != host:
        raise HTTPException(status_code=403, detail="invalid request origin")
    if request.headers.get("content-type", "").split(";", 1)[0] != "application/x-www-form-urlencoded":
        raise HTTPException(status_code=415, detail="form-encoded credentials are required")
    try:
        body = (await request.body()).decode("utf-8")
        form = parse_qs(body, keep_blank_values=True, strict_parsing=True)
        username = form.get("username", [""])[0]
        password = form.get("password", [""])[0]
    except (UnicodeDecodeError, ValueError):
        raise HTTPException(status_code=400, detail="invalid login form") from None
    if len(username) > 128 or len(password) > 1024:
        raise HTTPException(status_code=400, detail="invalid login form")

    client = request.client.host if request.client else "unknown"
    now = time.monotonic()
    for address, timestamps in list(login_failures.items()):
        recent = [timestamp for timestamp in timestamps if now - timestamp < login_window]
        if recent:
            login_failures[address] = recent
        else:
            login_failures.pop(address)
    recent_failures = [timestamp for timestamp in login_failures.get(client, []) if now - timestamp < login_window]
    login_failures[client] = recent_failures
    if len(recent_failures) >= login_failure_limit:
        raise HTTPException(status_code=429, detail="too many login attempts; try again later")

    try:
        _, _, serializer = auth_config()
        with closing(get_connection()) as connection:
            user = connection.execute(
                """
                SELECT id, username, password_hash, role, token_version
                FROM geocatalog_users
                WHERE username = %s AND is_active
                """,
                (username,),
            ).fetchone()
    except (RuntimeError, psycopg.Error) as exc:
        raise HTTPException(status_code=503, detail="authentication is not configured") from exc
    try:
        password_ok = user is not None and password_hasher.verify(user["password_hash"], password)
    except (VerifyMismatchError, InvalidHashError, VerificationError):
        password_ok = False
    if not password_ok:
        login_failures[client].append(now)
        response = RedirectResponse("/login?error=invalid", status_code=303)
        response.headers["Cache-Control"] = "no-store"
        return response

    login_failures.pop(client, None)
    response = RedirectResponse("/", status_code=303)
    response.set_cookie(
        "geocatalog_session",
        serializer.dumps({"id": user["id"], "version": user["token_version"]}),
        max_age=session_max_age,
        httponly=True,
        secure=True,
        samesite="strict",
        path="/",
    )
    return response


@app.post("/auth/logout")
def logout(request: Request):
    origin = request.headers.get("origin")
    if origin and urlsplit(origin).netloc != request.headers.get("host", ""):
        raise HTTPException(status_code=403, detail="invalid request origin")
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie("geocatalog_session", path="/", secure=True, httponly=True, samesite="strict")
    return response


@app.get("/healthz")
def healthz():
    try:
        with closing(get_connection()) as connection:
            connection.execute("SELECT PostGIS_Version()")
    except psycopg.Error as exc:
        raise HTTPException(status_code=503, detail="database unavailable") from exc
    return {"status": "ok"}


@app.get("/livez")
def livez():
    return {"status": "ok"}


@app.get("/api/features")
def search_features(
    q: Annotated[str | None, Query(max_length=200)] = None,
    region: Annotated[str | None, Query(max_length=200)] = None,
    year_from: Annotated[int | None, Query(ge=1500, le=2200)] = None,
    year_to: Annotated[int | None, Query(ge=1500, le=2200)] = None,
    bbox: Annotated[str | None, Query(max_length=100)] = None,
    limit: Annotated[int, Query(ge=1, le=500)] = 200,
):
    if year_from is not None and year_to is not None and year_from > year_to:
        raise HTTPException(status_code=422, detail="year_from must not exceed year_to")
    bounds = parse_bbox(bbox)
    conditions = []
    params = []
    if q:
        conditions.append(
            "(d.search_vector @@ websearch_to_tsquery('simple', %s) "
            "OR d.search_vector @@ websearch_to_tsquery('russian', %s) "
            "OR d.search_vector @@ websearch_to_tsquery('english', %s) "
            "OR d.authors ILIKE %s OR d.coauthors ILIKE %s OR d.executor_org ILIKE %s "
            "OR d.document_type ILIKE %s OR d.minerals ILIKE %s OR d.archive_disk_number ILIKE %s "
            "OR d.tgf_number ILIKE %s)"
        )
        params.extend((q, q, q, *(f"%{q}%" for _ in range(7))))
    if region:
        conditions.append("d.region ILIKE %s")
        params.append(f"%{region}%")
    if year_from is not None:
        conditions.append("coalesce(d.work_year_start, d.year) >= %s")
        params.append(year_from)
    if year_to is not None:
        conditions.append("coalesce(d.work_year_end, d.year) <= %s")
        params.append(year_to)
    if bounds:
        conditions.append("ST_Intersects(f.geom, ST_MakeEnvelope(%s, %s, %s, %s, 4326))")
        params.extend(bounds)
    where_clause = " AND ".join(conditions) if conditions else "TRUE"
    query = f"""
        SELECT f.id, f.name, f.kind, f.metadata, ST_AsGeoJSON(f.geom) AS geometry,
               ST_AsGeoJSON(f.source_geom) AS source_geometry, f.source_crs,
               json_agg({document_json_sql()}
                   ORDER BY d.year DESC NULLS LAST, d.title) AS documents
        FROM features f
        JOIN feature_documents fd ON fd.feature_id = f.id
        JOIN documents d ON d.id = fd.document_id
        WHERE {where_clause}
        GROUP BY f.id
        ORDER BY f.name
        LIMIT %s
    """
    with closing(get_connection()) as connection:
        rows = connection.execute(query, [*params, limit]).fetchall()
        attach_document_files(connection, rows)
    features = []
    for row in rows:
        features.append(
            {
                "type": "Feature",
                "id": row["id"],
                "geometry": decode_geojson(row["geometry"]),
                "properties": {
                    "name": row["name"],
                    "kind": row["kind"],
                    "metadata": row["metadata"],
                    "source_geometry": decode_geojson(row["source_geometry"]),
                    "source_crs": row["source_crs"],
                    "documents": row["documents"],
                },
            }
        )
    return {"type": "FeatureCollection", "features": features}


@app.get("/api/features/{feature_id}")
def get_feature(feature_id: int):
    with closing(get_connection()) as connection:
        row = connection.execute(
            f"""
            SELECT f.id, f.name, f.kind, f.metadata, ST_AsGeoJSON(f.geom) AS geometry,
                   ST_AsGeoJSON(f.source_geom) AS source_geometry, f.source_crs,
                   coalesce(json_agg({document_json_sql()}
                       ORDER BY d.year DESC NULLS LAST, d.title)
                   FILTER (WHERE d.id IS NOT NULL), '[]'::json) AS documents
            FROM features f
            LEFT JOIN feature_documents fd ON fd.feature_id = f.id
            LEFT JOIN documents d ON d.id = fd.document_id
            WHERE f.id = %s
            GROUP BY f.id
            """,
            (feature_id,),
        ).fetchone()
        if row is not None:
            attach_document_files(connection, [row])
    if row is None:
        raise HTTPException(status_code=404, detail="feature not found")
    return {
        "type": "Feature",
        "id": row["id"],
        "geometry": decode_geojson(row["geometry"]),
        "properties": {
            "name": row["name"],
            "kind": row["kind"],
            "metadata": row["metadata"],
            "source_geometry": decode_geojson(row["source_geometry"]),
            "source_crs": row["source_crs"],
            "documents": row["documents"],
        },
    }
